# -*- coding: utf-8 -*-
"""读取相关参考资料的工具脚本"""
import sys
import io
import os
import fitz  # PyMuPDF
import docx

sys.stdout.reconfigure(encoding='utf-8')

def read_docx(path):
    print(f"\n========== {path} ==========")
    try:
        doc = docx.Document(path)
        for p in doc.paragraphs:
            if p.text.strip():
                print(p.text)
        for tbl in doc.tables:
            for row in tbl.rows:
                cells = [c.text.strip() for c in row.cells]
                print(" | ".join(cells))
    except Exception as e:
        print(f"ERR: {e}")


def read_pdf_excerpt(path, max_pages=4):
    print(f"\n========== PDF: {os.path.basename(path)} ==========")
    try:
        doc = fitz.open(path)
        # 摘要、引言通常在前几页；实验通常在中段
        pages_to_read = list(range(min(max_pages, len(doc))))
        # 如果有更多页，采样实验/setup区域
        if len(doc) > 6:
            pages_to_read += [len(doc) // 2, len(doc) // 2 + 1]
        for i in pages_to_read:
            if i < len(doc):
                text = doc[i].get_text("text")
                # 截取前 1500 字
                print(f"--- page {i+1} ---")
                print(text[:1500])
        doc.close()
    except Exception as e:
        print(f"ERR: {e}")


if __name__ == "__main__":
    targets = sys.argv[1:]
    for t in targets:
        if t.lower().endswith('.docx'):
            read_docx(t)
        elif t.lower().endswith('.pdf'):
            read_pdf_excerpt(t)
